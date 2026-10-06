#!/usr/bin/env python3
"""Export this skill's LaTeX lab template to a DOCX for final visual review."""

import argparse
from copy import deepcopy
from datetime import date
import json
import re
from pathlib import Path
from os import pathsep
import subprocess
from tempfile import TemporaryDirectory

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def format_document(document, heading_space_pt=21, table_spacing=1, caption_spacing=1, figure_scale=1):
    for section in document.sections:
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin = section.bottom_margin = Cm(2)
        section.left_margin, section.right_margin = Cm(2.5), Cm(1.5)
        section.header_distance = Cm(0.8)
        section.different_first_page_header_footer = True
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header.paragraph_format.first_line_indent = Cm(0)
        field = OxmlElement('w:fldSimple')
        field.set(qn('w:instr'), 'PAGE')
        header._p.append(field)

    for style in document.styles:
        if style.type != WD_STYLE_TYPE.PARAGRAPH:
            continue
        style.font.name, style.font.size = 'Times New Roman', Pt(14)
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.rPr.rFonts
        for attribute in list(fonts.attrib):
            if 'theme' in attribute.lower():
                del fonts.attrib[attribute]
        fonts.set(qn('w:cs'), 'Times New Roman')
        fonts.set(qn('w:eastAsia'), 'Times New Roman')
        paragraph = style.paragraph_format
        paragraph.line_spacing = 1.5
        paragraph.space_before = paragraph.space_after = Pt(0)
        paragraph.widow_control = True

    for name in ('Normal', 'Body Text', 'First Paragraph'):
        paragraph = document.styles[name].paragraph_format
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.first_line_indent = Cm(1.25)

    for level in range(1, 10):
        style = document.styles[f'Heading {level}']
        style.font.bold = True
        paragraph = style.paragraph_format
        paragraph.keep_with_next = True
        paragraph.first_line_indent = Cm(0 if level == 1 else 1.25)
        paragraph.space_before = paragraph.space_after = Pt(heading_space_pt)
        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
        )

    for name in ('Source Code', 'Verbatim Char'):
        if name in document.styles:
            document.styles[name].font.name = 'Courier New'
            document.styles[name].font.size = Pt(12)

    document.styles['Hyperlink'].font.color.rgb = RGBColor(0, 0, 0)
    document.styles['Hyperlink'].font.underline = False
    figure_style = document.styles['Captioned Figure'].paragraph_format
    figure_style.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure_style.first_line_indent = Cm(0)
    figure_style.keep_with_next = True
    figure_style.line_spacing = caption_spacing

    title_paragraphs = []
    for paragraph in document.paragraphs:
        if paragraph.style.name.startswith('Heading'):
            first_heading = paragraph
            break
        title_paragraphs.append(paragraph)
    else:
        raise ValueError('Expected a title page followed by a section heading.')

    # Pandoc discards title spacing; split explicit lines to restore it.
    lines = []
    for paragraph in title_paragraphs:
        line = paragraph.insert_paragraph_before()
        for child in list(paragraph._p):
            if child.tag == qn('w:pPr'):
                continue
            if child.tag == qn('w:r') and child.find(qn('w:br')) is not None:
                lines.append(line)
                line = paragraph.insert_paragraph_before()
            else:
                line._p.append(deepcopy(child))
        lines.append(line)
        paragraph._p.getparent().remove(paragraph._p)

    right_block = False
    for index, paragraph in enumerate(lines):
        text = paragraph.text.strip()
        fmt = paragraph.paragraph_format
        fmt.first_line_indent = Cm(0)
        fmt.line_spacing = 1
        fmt.keep_with_next = True
        fmt.keep_together = True
        if text.startswith('Лектор:'):
            right_block = True
        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.RIGHT if right_block else WD_ALIGN_PARAGRAPH.CENTER
        )
        if index == 2:
            fmt.space_before = Pt(20)
        if text.startswith('Кафедра '):
            fmt.space_before = Pt(6)
        if paragraph._p.xpath('.//w:drawing'):
            fmt.space_before = Pt(10)
        if text == 'ЗВІТ':
            paragraph.style = document.styles['Title']
            fmt.space_before, fmt.space_after = Pt(20), Pt(10)
        if text.startswith('Лектор:'):
            fmt.space_before = Pt(24)
        elif text.startswith(('Виконав:', 'Виконала:', 'Прийняв:', 'Прийняла:', '«')):
            fmt.space_before = Pt(16)
    last = lines[-1]
    last.alignment = WD_ALIGN_PARAGRAPH.CENTER
    last.paragraph_format.space_before = Pt(20)
    last.paragraph_format.keep_with_next = False

    # A separate section preserves the body margins and page numbering.
    title_section = deepcopy(document.sections[0]._sectPr)
    last._p.get_or_add_pPr().append(title_section)
    margins = title_section.find(qn('w:pgMar'))
    for side, value in [('top', 1.5), ('bottom', 1.5), ('left', 2), ('right', 2)]:
        margins.set(qn(f'w:{side}'), str(Cm(value).twips))
    body_section = document.sections[-1]
    body_section.different_first_page_header_footer = False
    first_heading.paragraph_format.page_break_before = False

    for style_name, label in (('Image Caption', 'Рисунок'), ('Table Caption', 'Таблиця')):
        document.styles[style_name].font.italic = False
        captions = [p for p in document.paragraphs if p.style.name == style_name]
        for number, paragraph in enumerate(captions, 1):
            fmt = paragraph.paragraph_format
            fmt.first_line_indent = Cm(1.25 if label == 'Таблиця' else 0)
            fmt.keep_with_next = label == 'Таблиця'
            fmt.line_spacing = caption_spacing
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT if label == 'Таблиця'
                else WD_ALIGN_PARAGRAPH.CENTER
            )
            if not paragraph.text.startswith(label):
                suffix = '' if label == 'Таблиця' and len(captions) == 1 else f' {number}'
                run = paragraph.add_run(f'{label}{suffix} — ')
                paragraph._p.remove(run._r)
                paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, run._r)

    for table in document.tables:
        for index, row in enumerate(table.rows):
            properties = row._tr.get_or_add_trPr()
            properties.append(OxmlElement('w:cantSplit'))
            if index == 0:
                properties.append(OxmlElement('w:tblHeader'))
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    fmt = paragraph.paragraph_format
                    fmt.first_line_indent = Cm(0)
                    fmt.line_spacing = table_spacing
                    fmt.keep_with_next = (
                        index < len(table.rows) - 1 if len(table.rows) <= 10 else index == 0
                    )
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

    title_image_ids = {
        element.get('id') for paragraph in lines
        for element in paragraph._p.xpath('.//wp:docPr')
    }
    for shape in document.inline_shapes:
        if str(shape._inline.docPr.id) not in title_image_ids:
            shape.width = round(shape.width * figure_scale)
            shape.height = round(shape.height * figure_scale)

    bibliography_started = False
    for paragraph in document.paragraphs:
        if paragraph.style.name.startswith('Heading'):
            bibliography_started = paragraph.text == 'ПЕРЕЛІК ДЖЕРЕЛ ПОСИЛАННЯ'
        elif bibliography_started:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.keep_together = True


def prepare_bibliography(latex):
    """Keep the template's manual bibliography as numbered paragraphs."""
    pattern = r'\\begin\{thebibliography\}\{[^}]*\}(.*?)\\end\{thebibliography\}'
    matches = list(re.finditer(pattern, latex, re.S))
    if not matches:
        return latex, {}
    if len(matches) != 1:
        raise ValueError('Only one thebibliography block is supported.')
    block = matches[0]
    items = list(re.finditer(r'\\bibitem\{([^}]+)\}', block[1]))
    if not items or len(re.findall(r'\\bibitem\b', block[1])) != len(items):
        raise ValueError('Use plain \\bibitem{key}; optional labels are not supported.')
    numbers = {}
    entries = []
    for index, item in enumerate(items):
        key = item[1]
        if key in numbers:
            raise ValueError(f'Duplicate bibliography key: {key}')
        numbers[key] = index + 1
        end = items[index + 1].start() if index + 1 < len(items) else len(block[1])
        entry = block[1][item.end():end].strip()
        entries.append(r'{[' + str(index + 1) + ']} ' + entry + r'\par')
    replacement = r'\section*{ПЕРЕЛІК ДЖЕРЕЛ ПОСИЛАННЯ}' + '\n' + '\n'.join(entries)
    return latex[:block.start()] + replacement + latex[block.end():], numbers


def number_citations(node, numbers, used):
    """Use Pandoc's citation nodes, leaving code and comments untouched."""
    if isinstance(node, list):
        return [number_citations(item, numbers, used) for item in node]
    if not isinstance(node, dict):
        return node
    if node.get('t') == 'Cite':
        citations = node['c'][0]
        labels = []
        for citation in citations:
            key = citation['citationId']
            if key not in numbers:
                raise ValueError(f'Citation has no bibliography entry: {key}')
            if (citation['citationPrefix'] or citation['citationSuffix']
                    or citation['citationMode']['t'] != 'NormalCitation'):
                raise ValueError('Use plain \\cite{key}; citation notes/modes need separate conversion.')
            if key not in used:
                used.append(key)
            labels.append(str(numbers[key]))
        return {'t': 'Str', 'c': '[' + ', '.join(labels) + ']'}
    return {key: number_citations(value, numbers, used) for key, value in node.items()}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source', type=Path)
    parser.add_argument('output', type=Path)
    parser.add_argument('--pandoc', default='pandoc')
    parser.add_argument('--heading-space-pt', type=float, default=21)
    parser.add_argument('--table-line-spacing', type=float, default=1)
    parser.add_argument('--caption-line-spacing', type=float, default=1)
    parser.add_argument('--figure-scale', type=float, default=1)
    args = parser.parse_args()
    source, output = args.source.resolve(), args.output.resolve()
    if output.exists():
        parser.error(f'Output already exists: {output}')
    if source.suffix.lower() != '.tex' or output.suffix.lower() != '.docx':
        parser.error('Expected a .tex input and a .docx output.')
    if (args.heading_space_pt < 0 or args.table_line_spacing <= 0
            or args.caption_line_spacing <= 0 or not 0 < args.figure_scale <= 1):
        parser.error('Spacing must be positive (heading spacing may be zero); figure scale is 0 < scale <= 1.')
    with TemporaryDirectory(prefix='lpnu-docx-') as temp:
        intermediate = Path(temp) / 'converted.docx'
        # Keep title blanks inline instead of Pandoc block separators.
        latex = re.sub(r'\\rule\{([0-9.]+)cm\}\{0\.15mm\}',
                       lambda match: r'\_' * round(float(match[1]) * 4),
                       source.read_text().replace(r'\the\year', str(date.today().year)))
        resource_paths = [str(source.parent)]
        graphics = re.search(r'\\graphicspath\{((?:\{[^}]*\})+)\}', latex)
        if graphics:
            resource_paths.extend(str(source.parent / path)
                                  for path in re.findall(r'\{([^}]*)\}', graphics[1]))
        latex, numbers = prepare_bibliography(latex)
        parsed = subprocess.run([
            args.pandoc, '--from=latex', '--to=json', '--fail-if-warnings',
        ], cwd=source.parent, input=latex, text=True, capture_output=True, check=True)
        used = []
        ast = number_citations(json.loads(parsed.stdout), numbers, used)
        if list(numbers) != used:
            raise ValueError('Order bibitems by first citation and remove uncited entries in the LaTeX source.')
        subprocess.run([
            args.pandoc, '--from=json', '--to=docx',
            '--resource-path', pathsep.join(resource_paths),
            '--number-sections', '--metadata=lang:uk-UA', '--fail-if-warnings',
            '--output', str(intermediate),
        ], cwd=source.parent, input=json.dumps(ast), text=True, check=True)
        document = Document(intermediate)
        format_document(document, args.heading_space_pt, args.table_line_spacing,
                        args.caption_line_spacing, args.figure_scale)
        output.parent.mkdir(parents=True, exist_ok=True)
        document.save(output)
    print(f'Created {output}')
    print('Review text, captions, equation/citation references and every rendered page before delivery.')


if __name__ == '__main__':
    main()
